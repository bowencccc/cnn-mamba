#!/usr/bin/env python3
"""Create an English slide deck explaining the baseline and frame-aware model."""

from pathlib import Path
import tempfile

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Circle
import numpy as np
import pandas as pd
from PIL import Image, ImageOps, ImageDraw
from pptx import Presentation
from pptx.util import Inches
from sklearn.metrics import precision_recall_curve

from cnn_mamba.evaluate_candidates import reference_sites


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "presentations"
STEM = "cnn_mamba_architecture_evolution"

BG = "#F5F7FB"
INK = "#172A3A"
MUTED = "#607184"
GRID = "#DCE3EC"
BLUE = "#3478F6"
TEAL = "#159A8A"
ORANGE = "#F59E0B"
RED = "#E45756"
PURPLE = "#7C4DFF"
GREEN = "#54A24B"
TASKS = ("donor", "acceptor", "start", "stop")
TASK_TITLES = {
    "donor": "Donor", "acceptor": "Acceptor",
    "start": "Start codon", "stop": "Stop codon",
}
WHITE = "#FFFFFF"

plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 12,
    "axes.titleweight": "bold",
    "pdf.fonttype": 42,
})


def add_box(fig, x, y, w, h, text, color=WHITE, edge=GRID, fontsize=15,
            weight="normal", text_color=INK, radius=0.018, lw=1.4,
            align="center", zorder=2):
    patch = FancyBboxPatch(
        (x, y), w, h,
        boxstyle=f"round,pad=0.008,rounding_size={radius}",
        transform=fig.transFigure, facecolor=color, edgecolor=edge,
        linewidth=lw, zorder=zorder,
    )
    fig.add_artist(patch)
    tx = x + w / 2 if align == "center" else x + 0.018
    ha = "center" if align == "center" else "left"
    fig.text(tx, y + h / 2, text, ha=ha, va="center", fontsize=fontsize,
             weight=weight, color=text_color, zorder=zorder + 1,
             linespacing=1.25)
    return patch


def add_arrow(fig, x1, y1, x2, y2, color=MUTED, lw=2.0,
              style="-|>", connectionstyle="arc3"):
    arrow = FancyArrowPatch(
        (x1, y1), (x2, y2), transform=fig.transFigure,
        arrowstyle=style, mutation_scale=14, linewidth=lw,
        color=color, connectionstyle=connectionstyle, zorder=4,
    )
    fig.add_artist(arrow)
    return arrow


def add_pill(fig, x, y, text, color=BLUE, width=None, fontsize=10):
    width = width or max(0.07, 0.009 * len(text) + 0.025)
    add_box(fig, x, y, width, 0.035, text, color=color, edge=color,
            fontsize=fontsize, weight="bold", text_color=WHITE,
            radius=0.016, lw=0.0)


def add_bullets(fig, x, y, lines, fontsize=16, color=INK, gap=0.075,
                bullet_color=BLUE, width=0.41):
    current = y
    for line in lines:
        fig.add_artist(Circle((x, current + 0.004), 0.006,
                              transform=fig.transFigure,
                              facecolor=bullet_color, edgecolor="none", zorder=3))
        fig.text(x + 0.017, current, line, ha="left", va="center",
                 fontsize=fontsize, color=color, linespacing=1.25,
                 wrap=True, zorder=4)
        current -= gap


def add_image(fig, path, rect, border=True):
    ax = fig.add_axes(rect, zorder=2)
    image = plt.imread(path)
    ax.imshow(image)
    ax.axis("off")
    if border:
        for spine in ax.spines.values():
            spine.set_visible(True)
            spine.set_edgecolor(GRID)
            spine.set_linewidth(1.1)
    return ax


def base_slide(title, number, kicker=None, subtitle=None):
    fig = plt.figure(figsize=(13.333, 7.5), dpi=144, facecolor=BG)
    fig.add_artist(FancyBboxPatch(
        (0, 0.965), 1, 0.035, boxstyle="square,pad=0",
        transform=fig.transFigure, facecolor=BLUE, edgecolor="none",
    ))
    if kicker:
        fig.text(0.055, 0.913, kicker.upper(), fontsize=10.5,
                 weight="bold", color=BLUE, va="center")
        title_y = 0.855
    else:
        title_y = 0.875
    fig.text(0.055, title_y, title, fontsize=27, weight="bold",
             color=INK, va="center")
    if subtitle:
        fig.text(0.055, title_y - 0.052, subtitle, fontsize=12.5,
                 color=MUTED, va="center")
    fig.text(0.055, 0.028, "CNN–Mamba-k7 architecture evolution",
             fontsize=8.5, color=MUTED)
    fig.text(0.955, 0.028, f"{number:02d}", ha="right",
             fontsize=9, color=MUTED, weight="bold")
    return fig


def slide_title():
    fig = plt.figure(figsize=(13.333, 7.5), dpi=144, facecolor=INK)
    fig.add_artist(FancyBboxPatch(
        (0.055, 0.17), 0.012, 0.66, boxstyle="square,pad=0",
        transform=fig.transFigure, facecolor=ORANGE, edgecolor="none",
    ))
    fig.text(0.105, 0.69, "CNN–Mamba-k7", fontsize=42, weight="bold",
             color=WHITE, va="center")
    fig.text(0.105, 0.58, "From a two-head baseline to\nframe-aware auxiliary learning",
             fontsize=28, weight="bold", color="#DCE8F7", va="top",
             linespacing=1.2)
    fig.text(0.105, 0.36,
             "Drosophila • 10 kb windows • splice sites and translation boundaries",
             fontsize=15, color="#AFC2D9")
    add_pill(fig, 0.105, 0.25, "BASELINE", BLUE, 0.105, 10)
    add_arrow(fig, 0.22, 0.268, 0.34, 0.268, color="#86A2BF", lw=2.2)
    add_pill(fig, 0.355, 0.25, "DILATION-3 FRAME BRANCH", TEAL, 0.215, 10)
    add_arrow(fig, 0.585, 0.268, 0.705, 0.268, color="#86A2BF", lw=2.2)
    add_pill(fig, 0.72, 0.25, "PHASE AUXILIARY HEAD", PURPLE, 0.19, 10)
    fig.text(0.105, 0.10, "Architecture, supervision, inference, and measured impact",
             fontsize=12, color="#86A2BF")
    return fig


def slide_baseline_overview():
    fig = base_slide("Previous model: CNN–Mamba-k7 baseline", 2,
                     kicker="Architecture",
                     subtitle="One shared sequence encoder, two task-specific output heads")
    xs = [0.06, 0.22, 0.39, 0.58]
    widths = [0.115, 0.12, 0.16, 0.14]
    labels = [
        "DNA tokens\nA / C / G / T / N",
        "Embedding\n5 → 192",
        "Input projection\nLinear + GELU",
        "8 × CNN–BiMamba3\nblocks",
    ]
    colors = ["#EAF1FF", "#EAF1FF", "#EAF1FF", "#DDF4F0"]
    edges = [BLUE, BLUE, BLUE, TEAL]
    for x, w, label, color, edge in zip(xs, widths, labels, colors, edges):
        add_box(fig, x, 0.58, w, 0.14, label, color=color, edge=edge,
                fontsize=14, weight="bold")
    for x1, w1, x2 in zip(xs[:-1], widths[:-1], xs[1:]):
        add_arrow(fig, x1 + w1 + 0.006, 0.65, x2 - 0.008, 0.65)
    add_box(fig, 0.74, 0.58, 0.11, 0.14, "Final\nLayerNorm",
            color="#F0EDFF", edge=PURPLE, fontsize=14, weight="bold")
    add_arrow(fig, 0.726, 0.65, 0.742, 0.65)

    add_arrow(fig, 0.85, 0.65, 0.875, 0.73, color=BLUE,
              connectionstyle="arc3,rad=-0.15")
    add_arrow(fig, 0.85, 0.65, 0.875, 0.49, color=ORANGE,
              connectionstyle="arc3,rad=0.15")
    add_box(fig, 0.88, 0.67, 0.105, 0.15,
            "Splice head\n\nBG / donor /\nacceptor",
            color="#EAF1FF", edge=BLUE, fontsize=11.5, weight="bold")
    add_box(fig, 0.88, 0.38, 0.105, 0.15,
            "Start / stop head\n\nBG / start / stop",
            color="#FFF3DC", edge=ORANGE, fontsize=11.2, weight="bold")

    add_box(fig, 0.06, 0.19, 0.42, 0.23,
            "Training objective\n\nL = 0.25 × Lsplice  +  1.0 × Lstart/stop",
            color=WHITE, edge=GRID, fontsize=18, weight="bold")
    add_bullets(fig, 0.55, 0.36, [
        "10 kb input windows with 5 kb stride",
        "Canonical-candidate cross-entropy",
        "No explicit reading-frame target",
        "Start and stop must be learned indirectly from sequence context",
    ], fontsize=14.2, gap=0.062, bullet_color=ORANGE)
    return fig


def slide_old_block():
    fig = base_slide("Inside each baseline CNN–BiMamba3 block", 3,
                     kicker="Shared encoder",
                     subtitle="Local motif extraction and long-range bidirectional context")
    add_box(fig, 0.07, 0.58, 0.12, 0.11, "Input x",
            color="#EAF1FF", edge=BLUE, fontsize=17, weight="bold")
    add_arrow(fig, 0.19, 0.635, 0.245, 0.635)
    add_box(fig, 0.25, 0.54, 0.23, 0.19,
            "CNN residual path\nLayerNorm\nDepthwise Conv1D, k=7, d=1\nGELU → 1×1 Conv → Dropout",
            color="#FFF3DC", edge=ORANGE, fontsize=13.5, weight="bold")
    add_arrow(fig, 0.48, 0.635, 0.535, 0.635)
    add_box(fig, 0.54, 0.54, 0.23, 0.19,
            "Bidirectional Mamba3 path\nLayerNorm\nForward Mamba3\n+ reverse-complement direction",
            color="#DDF4F0", edge=TEAL, fontsize=13.5, weight="bold")
    add_arrow(fig, 0.77, 0.635, 0.825, 0.635)
    add_box(fig, 0.83, 0.58, 0.12, 0.11, "Output x′",
            color="#F0EDFF", edge=PURPLE, fontsize=17, weight="bold")
    fig.text(0.25, 0.47,
             "Residual updates preserve the original representation at both stages.",
             fontsize=12.5, color=MUTED)

    positions = np.arange(-3, 4)
    ax = fig.add_axes([0.08, 0.17, 0.43, 0.20])
    ax.scatter(positions, np.zeros_like(positions), s=340, color=ORANGE,
               edgecolors=WHITE, linewidths=2, zorder=3)
    for value in positions:
        ax.text(value, 0, f"{value:+d}" if value else "0", ha="center",
                va="center", fontsize=11, color=WHITE, weight="bold")
    ax.plot([-3, 3], [0, 0], color=ORANGE, linewidth=3, zorder=1)
    ax.set_xlim(-3.7, 3.7); ax.set_ylim(-0.8, 0.8); ax.axis("off")
    fig.text(0.08, 0.37, "k=7, dilation=1 samples adjacent positions",
             fontsize=15, weight="bold", color=INK)
    add_bullets(fig, 0.59, 0.34, [
        "Strong splice-site performance from local motifs + long context",
        "Adjacent CNN taps mix all three codon offsets",
        "No auxiliary pressure to represent coding phase",
    ], fontsize=14.5, gap=0.075, bullet_color=RED)
    return fig


def slide_current_architecture():
    fig = base_slide("Current model: explicit frame-aware auxiliary learning", 4,
                     kicker="Architecture update",
                     subtitle="The original backbone and splice path are preserved")
    add_box(fig, 0.06, 0.58, 0.16, 0.12, "DNA → embedding",
            color="#EAF1FF", edge=BLUE, fontsize=16, weight="bold")
    add_arrow(fig, 0.22, 0.64, 0.27, 0.64)
    add_box(fig, 0.28, 0.55, 0.20, 0.18,
            "8 × CNN–BiMamba3\n+ final LayerNorm\n\nx ∈ ℝᴮˣᴸˣ¹⁹²",
            color="#DDF4F0", edge=TEAL, fontsize=15, weight="bold")

    add_arrow(fig, 0.48, 0.64, 0.57, 0.73, color=BLUE,
              connectionstyle="arc3,rad=-0.13")
    add_box(fig, 0.58, 0.68, 0.16, 0.12, "Splice head",
            color="#EAF1FF", edge=BLUE, fontsize=16, weight="bold")
    add_box(fig, 0.79, 0.68, 0.16, 0.12,
            "background\ndonor / acceptor", color=WHITE, edge=BLUE,
            fontsize=13, weight="bold")
    add_arrow(fig, 0.74, 0.74, 0.785, 0.74, color=BLUE)

    add_arrow(fig, 0.48, 0.61, 0.56, 0.50, color=TEAL,
              connectionstyle="arc3,rad=0.12")
    add_box(fig, 0.57, 0.43, 0.20, 0.16,
            "FrameCNNBranch\nk=7, dilation=3\ndepthwise → 1×1 → residual",
            color="#DDF4F0", edge=TEAL, fontsize=14, weight="bold")
    add_arrow(fig, 0.77, 0.51, 0.815, 0.58, color=ORANGE,
              connectionstyle="arc3,rad=-0.12")
    add_arrow(fig, 0.77, 0.49, 0.815, 0.37, color=PURPLE,
              connectionstyle="arc3,rad=0.12")
    add_box(fig, 0.82, 0.53, 0.15, 0.11, "Start / stop head",
            color="#FFF3DC", edge=ORANGE, fontsize=14, weight="bold")
    add_box(fig, 0.82, 0.30, 0.15, 0.11, "Phase head\nLinear 192 → 3",
            color="#F0EDFF", edge=PURPLE, fontsize=14, weight="bold")

    add_box(fig, 0.07, 0.15, 0.43, 0.21,
            "Combined training objective\n\nL = 0.25Lsplice + Lstart/stop + λphase Lphase",
            color=WHITE, edge=GRID, fontsize=17.5, weight="bold")
    add_bullets(fig, 0.57, 0.25, [
        "Start/stop and phase share frame-aware features",
        "Splice predictions bypass the added frame branch",
        "Phase head is auxiliary: used for training, not required by UniAnn",
    ], fontsize=13.8, gap=0.063, bullet_color=GREEN)
    return fig


def slide_dilation():
    fig = base_slide("What dilation = 3 changes", 5,
                     kicker="Frame branch",
                     subtitle="Same output length; different sampling pattern")
    ax1 = fig.add_axes([0.08, 0.56, 0.84, 0.14])
    coords = np.arange(-9, 10)
    colors = ["#D5DCE6"] * len(coords)
    for p in range(-3, 4):
        colors[p + 9] = ORANGE
    ax1.scatter(coords, np.zeros_like(coords), s=230, color=colors,
                edgecolors=WHITE, linewidths=1.5)
    ax1.plot([-9, 9], [0, 0], color=GRID, linewidth=3, zorder=0)
    for p in range(-3, 4):
        ax1.text(p, 0, str(p), ha="center", va="center", fontsize=9.5,
                 color=WHITE, weight="bold")
    ax1.set_xlim(-10, 10); ax1.set_ylim(-0.8, 0.8); ax1.axis("off")
    fig.text(0.08, 0.73, "Baseline k=7, dilation=1", fontsize=17,
             weight="bold", color=INK)
    fig.text(0.70, 0.73, "effective receptive field: 7 bp",
             fontsize=12.5, color=MUTED)

    ax2 = fig.add_axes([0.08, 0.30, 0.84, 0.14])
    selected = np.arange(-9, 10, 3)
    colors2 = [TEAL if p != 0 else BLUE for p in selected]
    ax2.scatter(selected, np.zeros_like(selected), s=260, color=colors2,
                edgecolors=WHITE, linewidths=1.5)
    ax2.plot([-9, 9], [0, 0], color=TEAL, linewidth=3, zorder=0)
    for p in selected:
        ax2.text(p, 0, str(p), ha="center", va="center", fontsize=9.5,
                 color=WHITE, weight="bold")
    ax2.set_xlim(-10, 10); ax2.set_ylim(-0.8, 0.8); ax2.axis("off")
    fig.text(0.08, 0.47, "Frame branch k=7, dilation=3", fontsize=17,
             weight="bold", color=INK)
    fig.text(0.68, 0.47, "effective receptive field: 19 bp",
             fontsize=12.5, color=MUTED)

    add_box(fig, 0.08, 0.12, 0.84, 0.115,
            "Taps are spaced by 3 bp: i−9, i−6, i−3, i, i+3, i+6, i+9.\n"
            "They preserve the same codon offset without striding or downsampling.",
            color=WHITE, edge=TEAL, fontsize=13.2, weight="bold")
    return fig


def slide_phase_labels():
    fig = base_slide("From local GFF phase to per-base supervision", 6,
                     kicker="Training labels",
                     subtitle="The annotation is expanded; it is never supplied as model input")
    add_box(fig, 0.05, 0.60, 0.21, 0.15,
            "EviAnn CDS feature\nGFF column 8\nphase = 0 / 1 / 2",
            color="#FFF3DC", edge=ORANGE, fontsize=12.8, weight="bold")
    add_arrow(fig, 0.25, 0.675, 0.31, 0.675)
    add_box(fig, 0.32, 0.60, 0.21, 0.15,
            "Convert feature start\ninitial =\n(3 − phase) mod 3",
            color="#EAF1FF", edge=BLUE, fontsize=12.8, weight="bold")
    add_arrow(fig, 0.52, 0.675, 0.58, 0.675)
    add_box(fig, 0.59, 0.60, 0.35, 0.15,
            "Expand every CDS nucleotide\nphase[i] =\n(initial + transcript offset) mod 3",
            color="#F0EDFF", edge=PURPLE, fontsize=12.8, weight="bold")

    fig.text(0.07, 0.52, "Example", fontsize=16, weight="bold", color=INK)
    bases = list("ATGGAATAA")
    phases = [0, 1, 2, 0, 1, 2, 0, 1, 2]
    start_x, y = 0.10, 0.42
    for index, (base, phase) in enumerate(zip(bases, phases)):
        x = start_x + index * 0.065
        add_box(fig, x, y, 0.048, 0.065, base,
                color=WHITE, edge=GRID, fontsize=16, weight="bold", radius=0.008)
        add_box(fig, x, y - 0.075, 0.048, 0.052, str(phase),
                color=["#EAF1FF", "#DDF4F0", "#F0EDFF"][phase],
                edge=[BLUE, TEAL, PURPLE][phase], fontsize=14,
                weight="bold", radius=0.008)
    fig.text(0.10, 0.30, "per-base phase target", fontsize=11.5, color=MUTED)

    add_bullets(fig, 0.70, 0.48, [
        "Reverse strand is converted in transcript direction",
        "Outside CDS: masked",
        "Conflicting isoform phases: masked",
        "All 115,688 EviAnn CDS junctions passed frame-continuity checks",
    ], fontsize=13.2, gap=0.06, bullet_color=PURPLE)
    add_box(fig, 0.08, 0.12, 0.84, 0.115,
            "Each 10 kb sample stores phase_labels[10,000] and phase_mask[10,000].\n"
            "Only mask = 1 positions contribute to cross-entropy.",
            color=WHITE, edge=GRID, fontsize=12.8, weight="bold")
    return fig


def slide_training_inference():
    fig = base_slide("Training and inference are intentionally different", 7,
                     kicker="No annotation leakage",
                     subtitle="Phase annotation shapes the representation but is not required at prediction time")
    add_box(fig, 0.07, 0.54, 0.39, 0.25,
            "TRAINING\n\nDNA sequence → three outputs\n• splice logits\n• start/stop logits\n• phase logits\n\nPhase labels contribute λphase × CE",
            color="#EAF1FF", edge=BLUE, fontsize=16, weight="bold", align="left")
    add_box(fig, 0.54, 0.54, 0.39, 0.25,
            "INFERENCE\n\nDNA sequence → two outputs\n• splice probabilities\n• start/stop probabilities\n\nNo phase labels; phase logits are not requested",
            color="#DDF4F0", edge=TEAL, fontsize=16, weight="bold", align="left")
    add_arrow(fig, 0.46, 0.665, 0.53, 0.665, color=INK, lw=2.4)
    fig.text(0.495, 0.705, "trained weights", ha="center",
             fontsize=10.5, color=MUTED, weight="bold")

    add_box(fig, 0.09, 0.22, 0.34, 0.16,
            "Gradient path\nLphase → phase head → frame branch\n→ shared CNN–Mamba backbone",
            color=WHITE, edge=PURPLE, fontsize=15, weight="bold")
    add_box(fig, 0.57, 0.22, 0.34, 0.16,
            "Downstream compatibility\nSame six-column score file\nSame UniAnn interface and thresholds",
            color=WHITE, edge=ORANGE, fontsize=15, weight="bold")
    fig.text(0.50, 0.13,
             "The model learns frame-aware features; it does not receive the true frame at inference.",
             ha="center", fontsize=15, weight="bold", color=RED)
    return fig


def slide_ablation():
    fig = base_slide("Held-out chrX ablation: where the gain comes from", 8,
                     kicker="Evidence",
                     subtitle="Exact AUPRC over every canonical candidate on the complete chrX positive strand")
    columns = ["Model", "Donor", "Acceptor", "Start", "Stop", "Mean"]
    data = [
        ["Baseline", ".9472", ".9454", ".7176", ".5337", ".7860"],
        ["Dilation-3 only", ".9438", ".9417", ".7549", ".5888", ".8073"],
        ["Phase auxiliary", ".9516", ".9468", ".8412", ".8760", ".9039"],
        ["Combined", ".9559", ".9527", ".8510", ".8823", ".9105"],
    ]
    ax = fig.add_axes([0.055, 0.29, 0.56, 0.48])
    ax.axis("off")
    table = ax.table(cellText=data, colLabels=columns, cellLoc="center",
                     colLoc="center", loc="center", colWidths=[0.26, .14, .17, .14, .14, .14])
    table.auto_set_font_size(False); table.set_fontsize(11.5); table.scale(1, 2.25)
    for (row, col), cell in table.get_celld().items():
        cell.set_edgecolor(WHITE); cell.set_linewidth(2)
        if row == 0:
            cell.set_facecolor(INK); cell.get_text().set_color(WHITE)
            cell.get_text().set_weight("bold")
        elif row == 4:
            cell.set_facecolor("#DDF4F0"); cell.get_text().set_weight("bold")
        elif row % 2:
            cell.set_facecolor(WHITE)
        else:
            cell.set_facecolor("#EAF0F6")

    add_box(fig, 0.66, 0.58, 0.28, 0.15,
            "Phase supervision is the main driver\n\nStop: 0.5337 → 0.8760",
            color="#F0EDFF", edge=PURPLE, fontsize=17, weight="bold")
    add_box(fig, 0.66, 0.37, 0.28, 0.15,
            "Dilation-3 adds a smaller but\nconsistent frame-aware gain\n\nCombined mean: 0.9105",
            color="#DDF4F0", edge=TEAL, fontsize=15.5, weight="bold")
    add_box(fig, 0.66, 0.17, 0.28, 0.13,
            "No reference annotation from chrX\nwas used for training.",
            color=WHITE, edge=RED, fontsize=14, weight="bold")
    return fig


def _comparison_sources(regime):
    if regime == "heldout":
        baseline = (
            ROOT / "artifacts" / "scores" / "early_stopping_v1"
            / "heldout" / "w10"
        )
        current = (
            ROOT / "artifacts" / "scores" / "frame_ablation_v1"
            / "combined" / "heldout" / "w10"
        )
    else:
        baseline = (
            ROOT / "artifacts" / "scores" / "early_stopping_v1"
            / "chrxtrain" / "w10"
        )
        current = (
            ROOT / "artifacts" / "scores" / "frame_ablation_chrxtrain_v1"
            / "combined" / "chrxtrain" / "w10"
        )
    return baseline, current


def _thin_curve(recall, precision, max_points=6000):
    if len(recall) <= max_points:
        return recall, precision
    indices = np.linspace(0, len(recall) - 1, max_points, dtype=np.int64)
    return recall[indices], precision[indices]


def slide_four_task_comparison(regime, number):
    heldout = regime == "heldout"
    title = (
        "Four-task PR–sensitivity curves: held-out chrX"
        if heldout else
        "Four-task PR–sensitivity curves: chrX-in-training"
    )
    subtitle = (
        "Honest generalization: chrX excluded from model training"
        if heldout else
        "Diagnostic fit: chrX candidates come from a chromosome included in training"
    )
    fig = base_slide(title, number, kicker="Exact candidate evaluation",
                     subtitle=subtitle)
    truth, _, _ = reference_sites(
        "drosophila", ROOT / "data" / "raw" / "dmel_reference.gtf",
        ROOT / "data" / "raw" / "dmel_genome.fa", "NC_004354.4",
    )
    baseline, current = _comparison_sources(regime)
    baseline_metrics = {
        row.task: float(row.AUPRC)
        for row in pd.read_csv(
            baseline / "candidate_auprc.tsv", sep="\t"
        ).itertuples()
    }
    current_metrics = {
        row.task: float(row.AUPRC)
        for row in pd.read_csv(
            current / "candidate_auprc.tsv", sep="\t"
        ).itertuples()
    }
    grid = fig.add_gridspec(
        2, 2, left=0.065, right=0.965, bottom=0.095, top=0.76,
        hspace=0.34, wspace=0.23,
    )
    for axis, task in zip(
        [fig.add_subplot(grid[row, col]) for row in range(2) for col in range(2)],
        TASKS,
    ):
        expected_positions = None
        labels = None
        for directory, label, color, metrics in (
            (baseline, "Previous CNN–Mamba-k7", MUTED, baseline_metrics),
            (current, "Current frame model, phase=0.10", PURPLE, current_metrics),
        ):
            positions = np.load(directory / f"{task}_positions0.npy", mmap_mode="r")
            scores = np.load(directory / f"{task}_scores.npy", mmap_mode="r")
            if expected_positions is None:
                expected_positions = np.asarray(positions)
                labels = np.isin(positions, truth[task], assume_unique=True)
            elif not np.array_equal(positions, expected_positions):
                raise RuntimeError(f"candidate positions differ for {regime} {task}")
            precision, recall, _ = precision_recall_curve(labels, scores)
            recall, precision = _thin_curve(recall, precision)
            axis.plot(
                recall, precision, color=color, linewidth=2.2,
                label=f"{label}  (AP={metrics[task]:.4f})",
            )
        axis.set_xlim(0, 1); axis.set_ylim(0, 1.01)
        axis.set_title(TASK_TITLES[task], fontsize=14, weight="bold", color=INK)
        axis.set_xlabel("Sensitivity (recall)", fontsize=10.5)
        axis.set_ylabel("Precision", fontsize=10.5)
        axis.tick_params(labelsize=9)
        axis.grid(alpha=0.22)
        axis.legend(loc="lower left", fontsize=8.3, frameon=True, framealpha=0.94)
        axis.set_facecolor(WHITE)
    return fig


def slide_heldout_curves():
    return slide_four_task_comparison("heldout", 9)


def slide_chrxtrain_curves():
    return slide_four_task_comparison("chrxtrain", 10)


def slide_weight_grid():
    fig = base_slide("Phase-loss coefficient: no single value wins every metric", 11,
                     kicker="Hyperparameter grid",
                     subtitle="dilation=3 fixed; only λphase changes")
    plot_path = ROOT / "results" / "phase_weight_grid_v1" / "strict_locus_metrics.png"
    add_image(fig, plot_path, [0.48, 0.23, 0.49, 0.55], border=False)
    columns = ["λphase", "Held-out\nmean AP", "In-train\nmean AP", "UniAnn\nF1", "+ EviAnn\nF1"]
    data = [
        ["0.03", ".9056", ".9334", "79.18", "90.18"],
        ["0.05", ".9105", ".9272", "76.68", "89.26"],
        ["0.10", ".9105", ".9274", "77.15", "90.54"],
        ["0.20", ".9024", ".9100", "72.97", "88.08"],
    ]
    ax = fig.add_axes([0.05, 0.34, 0.39, 0.41]); ax.axis("off")
    table = ax.table(cellText=data, colLabels=columns, cellLoc="center",
                     colLoc="center", loc="center",
                     colWidths=[.16, .22, .22, .18, .20])
    table.auto_set_font_size(False); table.set_fontsize(10.8); table.scale(1, 2.35)
    for (row, col), cell in table.get_celld().items():
        cell.set_edgecolor(WHITE); cell.set_linewidth(2)
        if row == 0:
            cell.set_facecolor(INK); cell.get_text().set_color(WHITE)
            cell.get_text().set_weight("bold")
        elif row in (1, 3):
            cell.set_facecolor("#EAF1FF" if row == 1 else "#DDF4F0")
        else:
            cell.set_facecolor(WHITE)
    add_box(fig, 0.06, 0.135, 0.36, 0.15,
            "Pareto choices\n0.03: UniAnn-only\n0.05: held-out AP  |  0.10: combined locus F1",
            color=WHITE, edge=ORANGE, fontsize=11.8, weight="bold")
    fig.text(0.50, 0.16,
             "0.20 is consistently worse.\n0.05 and 0.10 held-out mean AP are effectively tied.",
             fontsize=11.5, color=MUTED, weight="bold", linespacing=1.35)
    return fig


def slide_end_to_end():
    fig = base_slide("End-to-end evaluation and practical takeaway", 12,
                     kicker="Pipeline",
                     subtitle="Model quality is evaluated at both candidate and gene-model levels")
    stages = [
        (0.055, "Train", "held-out chrX\nor chrX-in-training", BLUE),
        (0.245, "Score", "all chrX+ canonical\ncandidates", TEAL),
        (0.435, "Export", "donor / acceptor /\nstart / stop scores", PURPLE),
        (0.625, "UniAnn", "assemble coding\ngene models", ORANGE),
        (0.815, "Strict compare", "gffread -M +\ngffcompare", RED),
    ]
    for x, title, subtitle, color in stages:
        add_box(fig, x, 0.58, 0.14, 0.16,
                f"{title}\n\n{subtitle}", color=WHITE, edge=color,
                fontsize=13.2, weight="bold")
    for left, right in zip(stages[:-1], stages[1:]):
        add_arrow(fig, left[0] + 0.14, 0.66, right[0] - 0.008, 0.66)

    add_box(fig, 0.07, 0.28, 0.25, 0.18,
            "Best honest candidate metric\n\nλphase = 0.05\nheld-out mean AP = 0.9105",
            color="#EAF1FF", edge=BLUE, fontsize=16, weight="bold")
    add_box(fig, 0.375, 0.28, 0.25, 0.18,
            "Best UniAnn-only locus result\n\nλphase = 0.03\nSn / Pr / F1 = 80.3 / 78.1 / 79.18",
            color="#DDF4F0", edge=TEAL, fontsize=15, weight="bold")
    add_box(fig, 0.68, 0.28, 0.25, 0.18,
            "Best EviAnn + UniAnn result\n\nλphase = 0.10\nSn / Pr / F1 = 91.6 / 89.5 / 90.54",
            color="#FFF3DC", edge=ORANGE, fontsize=15, weight="bold")
    fig.text(0.50, 0.16,
             "Recommended name: CNN–Mamba-k7 with a dilation-3 frame branch and phase auxiliary supervision",
             ha="center", fontsize=15.5, weight="bold", color=INK)
    fig.text(0.50, 0.105,
             "The architecture is fixed across the coefficient grid; only the phase-loss weight changes.",
             ha="center", fontsize=12.5, color=MUTED)
    return fig


def make_contact_sheet(paths, output):
    thumbs = []
    for path in paths:
        image = Image.open(path).convert("RGB")
        image.thumbnail((640, 360), Image.Resampling.LANCZOS)
        canvas = Image.new("RGB", (660, 390), "white")
        canvas.paste(image, ((660 - image.width) // 2, 10))
        thumbs.append(canvas)
    rows = (len(thumbs) + 1) // 2
    sheet = Image.new("RGB", (1320, 390 * rows), "#DCE3EC")
    for index, thumb in enumerate(thumbs):
        sheet.paste(thumb, ((index % 2) * 660, (index // 2) * 390))
    sheet.save(output, quality=92)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    pptx_path = OUT / f"{STEM}.pptx"
    pdf_path = OUT / f"{STEM}.pdf"
    preview_path = OUT / f"{STEM}_contact_sheet.png"

    builders = [
        slide_title,
        slide_baseline_overview,
        slide_old_block,
        slide_current_architecture,
        slide_dilation,
        slide_phase_labels,
        slide_training_inference,
        slide_ablation,
        slide_heldout_curves,
        slide_chrxtrain_curves,
        slide_weight_grid,
        slide_end_to_end,
    ]

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = (
        "CNN–Mamba-k7: from baseline to frame-aware auxiliary learning"
    )
    presentation.core_properties.subject = "Drosophila gene-signal model architecture"
    presentation.core_properties.author = "CNN–Mamba project"

    with tempfile.TemporaryDirectory(prefix="cnn_mamba_ppt_") as temp_name:
        temp = Path(temp_name)
        pngs = []
        with PdfPages(pdf_path) as pdf:
            for index, builder in enumerate(builders, 1):
                fig = builder()
                png = temp / f"slide_{index:02d}.png"
                fig.savefig(png, dpi=144, facecolor=fig.get_facecolor(),
                            edgecolor="none")
                pdf.savefig(fig, facecolor=fig.get_facecolor(), edgecolor="none")
                plt.close(fig)
                pngs.append(png)

                slide = presentation.slides.add_slide(
                    presentation.slide_layouts[6]
                )
                slide.shapes.add_picture(
                    str(png), 0, 0,
                    width=presentation.slide_width,
                    height=presentation.slide_height,
                )
        presentation.save(pptx_path)
        make_contact_sheet(pngs, preview_path)

    print(pptx_path)
    print(pdf_path)
    print(preview_path)


if __name__ == "__main__":
    main()
